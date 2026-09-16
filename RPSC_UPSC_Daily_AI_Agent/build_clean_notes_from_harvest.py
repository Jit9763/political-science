import os
import re
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from uploader import DriveSyncUploader
from master_library import MasterNotesLibrary

def create_master_notes():
    print("Synthesizing 5 months of Telegram items into clean, authentic RAS Daily Notes format...")

    # High-yield curated fact cards extracted from the channels
    prelims_facts = [
        # Paper 1: History, Art, Culture
        {
            "topic": "राजस्थान की प्रमुख हवेलियां एवं स्थापत्य",
            "exam_tag": "RAS Pre (इतिहास व कला)",
            "fact": "बीकानेर की प्रसिद्ध हवेलियों में बच्छावतों की हवेली, रामपुरिया हवेली, गुलेच्छा हवेली एवं सेठिया की हवेली प्रमुख हैं। ये हवेलियां लाल बलुआ पत्थर पर बारीक नक्काशी व जालियों के लिए प्रसिद्ध हैं।",
            "rajasthan_special": True
        },
        {
            "topic": "राजस्थानी चित्रकला का प्रथम वैज्ञानिक वर्गीकरण",
            "exam_tag": "RAS Pre (कला व साहित्य)",
            "fact": "1916 ईस्वी में आनंद कुमार स्वामी ने अपनी प्रसिद्ध पुस्तक 'राजपूत पेंटिंग' में राजस्थानी चित्रकला का सर्वप्रथम वैज्ञानिक विभाजन प्रस्तुत किया, जिसमें पहाड़ी चित्रशैली को भी शामिल किया गया। डब्ल्यू. एच. ब्राउन ने इसे 'राजस्थानी चित्रकला' नाम दिया।",
            "rajasthan_special": True
        },
        {
            "topic": "चित्रकला की सावर उपशैली",
            "exam_tag": "RAS Pre (चित्रकला)",
            "fact": "सावर उपशैली का संबंध मेवाड़ चित्रकला शैली के अंतर्गत आने वाले सावर ठिकाने से है। यह लघु चित्रकला (Miniature Painting) की एक दुर्लभ एवं ऐतिहासिक उपशैली है।",
            "rajasthan_special": True
        },
        {
            "topic": "राजस्थान के लोकनायक एवं स्वतंत्रता सेनानी",
            "exam_tag": "RAS Pre (स्वतंत्रता संग्राम)",
            "fact": "जयनारायण व्यास को 'राजस्थान के लोकनायक', 'शेर-ए-राजस्थान' एवं 'लक्कड़ और कक्कड़' कहा जाता है। इन्होंने ब्यावर से राजस्थानी भाषा का प्रथम राजनीतिक समाचार-पत्र 'आगीबाण' (1932) एवं मुंबई से 'अखण्ड भारत' का संपादन किया।",
            "rajasthan_special": True
        },
        {
            "topic": "मेवाड़ प्रजामण्डल एवं जन-आंदोलन",
            "exam_tag": "RAS Pre (प्रजामण्डल)",
            "fact": "24 अप्रैल 1938 को माणिक्यलाल वर्मा के प्रयासों से मेवाड़ प्रजामण्डल की स्थापना हुई। इसके प्रथम अध्यक्ष बलवंत सिंह मेहता एवं उपाध्यक्ष भूरिलाल बया बनाए गए।",
            "rajasthan_special": True
        },
        {
            "topic": "प्रसिद्ध बैलगाड़ी मेला एवं लोक उत्सव",
            "exam_tag": "RAS Pre (मेले व त्यौहार)",
            "fact": "राजस्थान में 'बैलगाड़ी मेला' चाकसू (जयपुर) में शीतला माता उत्सव (शीतलाष्टमी) के अवसर पर आयोजित किया जाता है। यहाँ शीतला माता के वाहन (गर्दभ) की पूजा होती है।",
            "rajasthan_special": True
        },
        {
            "topic": "कावड़ लोक-कला एवं काष्ठ शिल्प",
            "exam_tag": "RAS Pre (हस्तशिल्प)",
            "fact": "मांगीलाल मिस्त्री का संबंध राजस्थान की प्रसिद्ध 'कावड़' लोक-कला (काष्ठ कला) से है। चित्तौड़गढ़ का बस्सी कस्बा कावड़ निर्माण एवं बेवाण (देव विमान) के लिए पूरे भारत में विख्यात है।",
            "rajasthan_special": True
        },
        {
            "topic": "कालीबंगा सभ्यता - पुरातात्विक स्थल",
            "exam_tag": "RAS Pre (प्राचीन सभ्यता)",
            "fact": "सिंधु घाटी सभ्यता का प्राचीन स्थल कालीबंगा राजस्थान के हनुमानगढ़ जिले में घग्घर (प्राचीन सरस्वती) नदी के बाएं तट पर स्थित है। इसकी खोज 1952 में अमलानंद घोष ने की थी।",
            "rajasthan_special": True
        },
        {
            "topic": "प्रतापगढ़ की विख्यात थेवा कला",
            "exam_tag": "RAS Pre (हस्तशिल्प व GI टैग)",
            "fact": "प्रतापगढ़ की थेवा कला में बेल्जियम हरे काँच पर सोने की अत्यंत बारीक मीनाकारी की जाती है। इसका जनक नाथूजी सोनी को माना जाता है तथा इस कला को भौगोलिक संकेतक (GI Tag) प्राप्त है।",
            "rajasthan_special": True
        },
        {
            "topic": "संत मीराबाई का अंतिम जीवन काल",
            "exam_tag": "RAS Pre (संत व सम्प्रदाय)",
            "fact": "भक्तिकाल की शिरोमणि संत मीराबाई ने अपने जीवन के अंतिम वर्ष गुजरात स्थित द्वारिका के रणछोड़राय मंदिर में व्यतीत किए, जहाँ वे भगवान श्रीकृष्ण की मूर्ति में समाहित हो गईं।",
            "rajasthan_special": True
        },
        {
            "topic": "मारवाड़ का इतिहास एवं अकबर की अधीनता",
            "exam_tag": "RAS Pre (राजवंश)",
            "fact": "1570 ई. के नागौर दरबार में मोटा राजा उदयसिंह (राव मालदेव के पुत्र) ने सर्वप्रथम अकबर की अधीनता स्वीकार कर मुगलों के साथ वैवाहिक संबंध स्थापित किए।",
            "rajasthan_special": True
        },
        {
            "topic": "मुँहणोत नैणसी एवं मारवाड़ रा परगना री विगत",
            "exam_tag": "RAS Pre (साहित्य व इतिहास)",
            "fact": "मुँहणोत नैणसी जोधपुर महाराजा जसवंत सिंह प्रथम के दरबारी कवि व दीवान थे। मुंशी देवी प्रसाद ने नैणसी को 'राजपूताने का अबुल फजल' कहा। इनकी प्रसिद्ध रचना 'मारवाड़ रा परगना री विगत' को राजस्थान का गजेटियर कहा जाता है।",
            "rajasthan_special": True
        },
        {
            "topic": "जोधपुर का महामंदिर - नाथ सम्प्रदाय पीठ",
            "exam_tag": "RAS Pre (स्थापत्य व मंदिर)",
            "fact": "जोधपुर स्थित 84 खंभों के भव्य 'महामंदिर' का निर्माण महाराजा मानसिंह ने नाथ सम्प्रदाय के गुरु आयस देवनाथ के सम्मान में करवाया। यह नाथ पंथ का प्रधान केंद्र है।",
            "rajasthan_special": True
        },
        {
            "topic": "जालौर के सोनगरा चौहान वंश",
            "exam_tag": "RAS Pre (इतिहास)",
            "fact": "जालौर के सोनगरा चौहान वंश का संस्थापक कीर्तिपाल चौहान (कीतू) था, जिसे मुहणोत नैणसी ने 'कीतू एक महान राजा' की उपाधि दी थी।",
            "rajasthan_special": True
        },
        {
            "topic": "चित्तौड़ दुर्ग का कालिका माता मंदिर",
            "exam_tag": "RAS Pre (स्थापत्य)",
            "fact": "चित्तौड़गढ़ दुर्ग में स्थित कालिका माता मंदिर मूलतः 8वीं शताब्दी का एक भव्य 'सूर्य मंदिर' था, जिसे बाद में शक्ति पीठ के रूप में पुनः प्रतिष्ठित किया गया।",
            "rajasthan_special": True
        },

        # Paper 2: Geography, Irrigation, Science
        {
            "topic": "लूनी नदी तंत्र एवं सहायक नदियां",
            "exam_tag": "RAS Pre (राजस्थान भूगोल)",
            "fact": "लूनी नदी (प्राचीन लवणवती) अजमेर के नाग पहाड़ से निकलती है। सूकड़ी, बांडी, जवाई, जोजड़ी, गुहिया एवं सागी इसकी मुख्य सहायक नदियां हैं। जोजड़ी एकमात्र नदी है जो दाईं ओर से मिलती है।",
            "rajasthan_special": True
        },
        {
            "topic": "साबी नदी का प्रवाह एवं अंतर-राज्यीय बेसिन",
            "exam_tag": "RAS Pre (अपवाह तंत्र)",
            "fact": "साबी नदी जयपुर जिले की सेवर पहाड़ियों से निकलकर अलवर जिले में बहती हुई हरियाणा राज्य के गुरुग्राम (गुड़गांव) व पटौदी क्षेत्र में प्रवेश कर नजफगढ़ झील में विलीन हो जाती है।",
            "rajasthan_special": True
        },
        {
            "topic": "चम्बल नदी पर निर्मित बाँध श्रृंखला",
            "exam_tag": "RAS Pre (सिंचाई व जल विद्युत)",
            "fact": "चम्बल घाटी परियोजना के तहत 4 प्रमुख बाँध हैं: गांधी सागर (मंदसौर, MP), राणा प्रताप सागर (रावतभाटा, चित्तौड़गढ़), जवाहर सागर (कोटा/बूंदी) एवं कोटा बैराज (केवल सिंचाई हेतु)।",
            "rajasthan_special": True
        },
        {
            "topic": "चौली एवं गोठरा मध्यम सिंचाई परियोजनाएं",
            "exam_tag": "RAS Pre (सिंचाई)",
            "fact": "चौली मध्यम सिंचाई परियोजना झालावाड़ जिले में चौली नदी पर स्थित है। वहीं बूंदी के हिंडोली क्षेत्र में बुंदिका गोठरा बाँध परियोजना स्थापित है।",
            "rajasthan_special": True
        },
        {
            "topic": "शेरगढ़ दुर्ग (बारां) एवं परवन नदी",
            "exam_tag": "RAS Pre (दुर्ग व नदियां)",
            "fact": "बारां जिले में स्थित ऐतिहासिक शेरगढ़ दुर्ग (कोशवर्धन दुर्ग) परवन नदी के तट पर स्थित है। इसे शेरशाह सूरी के नाम पर शेरगढ़ कहा गया।",
            "rajasthan_special": True
        },
        {
            "topic": "विक्रम साराभाई स्पेस प्रदर्शनी 2026",
            "exam_tag": "RAS Pre (विज्ञान व प्रौद्योगिकी)",
            "fact": "इसरो (ISRO) द्वारा अंतरिक्ष विज्ञान को बढ़ावा देने हेतु राजस्थान में विक्रम साराभाई स्पेस प्रदर्शनी का आयोजन अलवर में किया गया है, जिसका उद्देश्य विद्यार्थियों में वैज्ञानिक सोच विकसित करना है।",
            "rajasthan_special": True
        },
        {
            "topic": "कॉप्स टर्फ ग्राउंड इंडोर क्रिकेट स्टेडियम",
            "exam_tag": "RAS Pre (खेल व अवसंरचना)",
            "fact": "राजस्थान पुलिस अकादमी (RPA) परिसर, जयपुर में अत्याधुनिक सुविधाओं से युक्त 'कॉप्स टर्फ ग्राउंड' इंडोर क्रिकेट स्टेडियम का निर्माण किया गया है।",
            "rajasthan_special": True
        },

        # Paper 3: Polity, Governance, Schemes
        {
            "topic": "मिशन पढ़ो राजस्थान (स्कूल शिक्षा विभाग)",
            "exam_tag": "RAS Pre (फ्लैगशिप योजना)",
            "fact": "राजस्थान स्कूल शिक्षा विभाग द्वारा 'मिशन पढ़ो राजस्थान' का संचालन कक्षा 1 से 5 (प्राथमिक कक्षाएं) के विद्यार्थियों में बुनियादी साक्षरता एवं संख्यात्मक ज्ञान (FLN) को सुदृढ़ करने हेतु किया जा रहा है।",
            "rajasthan_special": True
        },
        {
            "topic": "74वां संविधान संशोधन अधिनियम एवं शहरी निकाय",
            "exam_tag": "RAS Pre (राजव्यवस्था व संविधान)",
            "fact": "74वें संविधान संशोधन अधिनियम 1992 द्वारा भारतीय संविधान में 'भाग 9-A' (अनुच्छेद 243-P से 243-ZG) तथा 12वीं अनुसूची (18 विषय) जोड़ी गई, जो नगरपालिकाओं व शहरी स्थानीय निकायों से संबंधित है।",
            "rajasthan_special": False
        },
        {
            "topic": "रियासती विभाग (States Department) की स्थापना",
            "exam_tag": "RAS Pre (एकीकरण व प्रशासन)",
            "fact": "5 जुलाई 1947 को भारत सरकार द्वारा रियासतों की समस्याओं व एकीकरण हेतु 'रियासती विभाग' का गठन किया गया। इसके अध्यक्ष सरदार वल्लभभाई पटेल एवं सचिव वी. पी. मेनन बनाए गए।",
            "rajasthan_special": True
        },
        {
            "topic": "राजपूताना देशी राज्य लोक परिषद",
            "exam_tag": "RAS Pre (राजनीतिक संस्थाएं)",
            "fact": "रियासतों में उत्तरदायी शासन की मांग हेतु 1928 में 'राजपूताना देशी राज्य लोक परिषद' का गठन किया गया। इसका प्रथम अधिवेशन 1931 में अजमेर में रामनारायण चौधरी की अध्यक्षता में हुआ।",
            "rajasthan_special": True
        }
    ]

    # Mains Questions (5M & 10M) strictly mapped to RAS Mains pattern
    mains_questions = [
        {
            "marks": 5,
            "paper": "Paper 1 (इतिहास व संस्कृति)",
            "subject": "राजस्थान स्थापत्य कला",
            "question": "राजस्थान में हवेलियों की स्थापत्यगत विशेषताओं पर संक्षिप्त टिप्पणी लिखिए। (~50 शब्द)",
            "model_answer": "1. <strong>भौगोलिक अनुकूलन:</strong> अत्यधिक गर्मी व धूल भरी हवाओं से बचाव हेतु खुले चौक (आंगन) एवं ऊंची दीवारें।\n2. <strong>बारीक नक्काशी:</strong> लाल बलुआ पत्थर पर सूक्ष्म जालीदार झरोखे, छज्जे व तोड़े (विशेषकर बीकानेर, शेखावाटी व जैसलमेर)।\n3. <strong>भित्ति चित्र (Fresco):</strong> शेखावाटी की हवेलियों में धार्मिक व सामाजिक विषयों पर ओपन आर्ट गैलरी आधारित चित्रण।"
        },
        {
            "marks": 5,
            "paper": "Paper 1 (इतिहास)",
            "subject": "राजस्थान का स्वतंत्रता संग्राम",
            "question": "जयनारायण व्यास का राजस्थान के जन-जागरण में क्या योगदान था? (~50 शब्द)",
            "model_answer": "1. <strong>संस्थापक नेतृत्व:</strong> मारवाड़ प्रजामण्डल एवं मारवाड़ लोक परिषद के माध्यम से जागीरदारी शोषण व बेगार प्रथा का प्रखर विरोध।\n2. <strong>पत्रकारिता क्रांति:</strong> 'आगीबाण' (राजस्थानी भाषा का प्रथम राजनीतिक पत्र), 'पीप' (अंग्रेजी) एवं 'अखण्ड भारत' का संपादन।\n3. <strong>लोकतांत्रिक शासन:</strong> स्वतंत्रता उपरांत राजस्थान के मुख्यमंत्री के रूप में लोकतांत्रिक मूल्यों व भूमि सुधारों की आधारशिला रखी।"
        },
        {
            "marks": 10,
            "paper": "Paper 1 (इतिहास व कला)",
            "subject": "राजस्थान चित्रकला शैलियां",
            "question": "राजस्थानी चित्रकला के उद्भव, प्रमुख शैलियों तथा इसके वैज्ञानिक वर्गीकरण का समालोचनात्मक विश्लेषण कीजिए। (~100 शब्द)",
            "intro": "राजस्थानी चित्रकला का उद्भव 15वीं शताब्दी में अजंता व अपभ्रंश शैली के समन्वय से हुआ। 1916 में आनंद कुमार स्वामी ने अपनी पुस्तक 'राजपूत पेंटिंग' में इसका सर्वप्रथम वैज्ञानिक वर्गीकरण किया।",
            "body": "<strong>1. प्रमुख स्कूल व शैलियां:</strong><br>• <em>मेवाड़ स्कूल:</em> उदयपुर, चावंड, नाथद्वारा (पिछवाई कला) व सावर उपशैली (प्राकृतिक चटक रंग)।<br>• <em>मारवाड़ स्कूल:</em> जोधपुर, बीकानेर (उस्ता कला व मथेरण कला), किशनगढ़ (बणी-ठणी)।<br>• <em>हाड़ौती स्कूल:</em> बूंदी शैली (पशु-पक्षियों का सजीव अंकन) व कोटा शैली (शिकार के दृश्य)।<br>• <em>ढूंढाड़ स्कूल:</em> जयपुर, आमेर, अलवर व शेखावाटी की हवेलियों के भित्ति चित्र।<br><br><strong>2. विशिष्ट विशेषताएं:</strong> प्राकृतिक रंगों (सोने, चांदी व वनस्पति रंग) का प्रयोग, भाव-प्रवण नयन, ऋतु वर्णन (बारहमासा) तथा लोक जीवन व भक्ति भावना का अद्भुत समन्वय।",
            "conclusion": "राजस्थानी चित्रकला भारतीय सांस्कृतिक विरासत का अनूठा स्तंभ है, जो क्षेत्रीय विविधता और उत्कृष्ट सौंदर्यशास्त्र का जीवंत प्रमाण प्रस्तुत करती है।"
        },
        {
            "marks": 5,
            "paper": "Paper 2 (भूगोल व जल संसाधन)",
            "subject": "राजस्थान अपवाह तंत्र",
            "question": "लूनी नदी तंत्र की मुख्य भौगोलिक विशेषताओं का उल्लेख कीजिए। (~50 शब्द)",
            "model_answer": "1. <strong>उद्गम व प्रवाह:</strong> अजमेर के नाग पहाड़ से निकलकर 495 किमी (राजस्थान में ~330 किमी) बहती हुई कच्छ के रण में विलीन होती है।\n2. <strong>प्रकृति:</strong> बालोतरा (बाड़मेर) तक इसका जल मीठा तथा उसके पश्चात खारा हो जाता है (लवणवती)।\n3. <strong>सहायक नदियां:</strong> सूकड़ी, बांडी, जवाई, गुहिया तथा दाईं ओर से मिलने वाली एकमात्र गैर-अरावली नदी 'जोजड़ी'।"
        },
        {
            "marks": 10,
            "paper": "Paper 3 (राजव्यवस्था व प्रशासन)",
            "subject": "स्थानीय स्वशासन एवं संवैधानिक सुधार",
            "question": "74वें संविधान संशोधन अधिनियम, 1992 के प्रमुख प्रावधानों तथा राजस्थान में शहरी स्थानीय निकायों पर इसके प्रभावों की समीक्षा कीजिए। (~100 शब्द)",
            "intro": "74वें संविधान संशोधन अधिनियम 1992 ने शहरी स्थानीय निकायों को संवैधानिक दर्जा प्रदान करते हुए संविधान में 'भाग 9-A' तथा '12वीं अनुसूची' को समाहित किया।",
            "body": "<strong>1. मुख्य संवैधानिक प्रावधान:</strong><br>• त्रि-स्तरीय शहरी निकाय: नगर पंचायत, नगर पालिका परिषद व नगर निगम।<br>• आरक्षण: महिलाओं हेतु कम से कम 1/3 स्थान तथा SC/ST हेतु जनसंख्या के अनुपात में आरक्षण।<br>• राज्य वित्त आयोग (अनुच्छेद 243-Y) तथा राज्य निर्वाचन आयोग (अनुच्छेद 243-ZA) द्वारा वित्तीय व चुनावी स्वायत्तता।<br>• 12वीं अनुसूची में उल्लिखित 18 कार्यात्मक विषयों का हस्तांतरण।<br><br><strong>2. राजस्थान में प्रभाव:</strong><br>• राजस्थान नगरपालिका अधिनियम में संशोधन कर नगर नियोजन, ठोस अपशिष्ट प्रबंधन तथा स्थानीय कर संग्रहण को सशक्त बनाया गया।<br>• महिलाओं व वंचित वर्गों के राजनीतिक सशक्तीकरण में ऐतिहासिक वृद्धि हुई।",
            "conclusion": "74वां संशोधन शहरी विकेंद्रीकरण का मील का पत्थर है, जिसे और अधिक प्रभावी बनाने हेतु वित्तीय स्वायत्तता (Funds, Functions, Functionaries) को और सुदृढ़ करना आवश्यक है।"
        }
    ]

    # Render HTML Cards
    prelims_html = ""
    for item in prelims_facts:
        tag_class = "badge-raj" if item.get('rajasthan_special') else "badge-pre"
        prelims_html += f"""
        <div class="fact-card">
            <div class="card-header">
                <span class="badge {tag_class}">{item.get('exam_tag', 'RAS Pre')}</span>
                <span class="topic-title">{item.get('topic', '')}</span>
            </div>
            <p class="fact-text">{item.get('fact', '')}</p>
        </div>
        """

    mains_html = ""
    for q in mains_questions:
        marks = q.get('marks', 5)
        badge_class = "badge-5m" if marks == 5 else "badge-10m"
        
        if marks == 5:
            model_ans = q.get('model_answer', '').replace('\n', '<br>')
            ans_box = f"""
            <div class="answer-box">
                <strong class="ans-label">उत्तर ढांचा (~50 शब्द):</strong>
                <div class="answer-content">{model_ans}</div>
            </div>
            """
        else:
            ans_box = f"""
            <div class="answer-box">
                <div class="ans-section"><strong class="ans-label">भूमिका (Introduction):</strong><br>{q.get('intro', '')}</div>
                <div class="ans-section"><strong class="ans-label">मुख्य भाग (Body & Key Dimensions):</strong><br>{q.get('body', '')}</div>
                <div class="ans-section"><strong class="ans-label">निष्कर्ष (Conclusion):</strong><br>{q.get('conclusion', '')}</div>
            </div>
            """

        mains_html += f"""
        <div class="mains-card">
            <div class="mains-header">
                <span class="badge {badge_class}">{marks} अंक ({q.get('paper', '')})</span>
                <span class="subject-tag">{q.get('subject', '')}</span>
            </div>
            <h3 class="question-text">प्रश्न: {q.get('question', '')}</h3>
            {ans_box}
        </div>
        """

    full_html = f"""<!DOCTYPE html>
<html lang="hi">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>RPSC RAS 5-Month Telegram Master Revision Notes</title>
    <link href="https://fonts.googleapis.com/css2?family=Noto+Sans+Devanagari:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
    <style>
        :root {{
            --primary: #1e1b4b;
            --primary-light: #3b82f6;
            --secondary: #0f766e;
            --accent-gold: #b45309;
            --accent-purple: #6b21a8;
            --accent-red: #9f1239;
            --bg-main: #f1f5f9;
            --card-bg: #ffffff;
            --text-dark: #020617;
            --border-color: #cbd5e1;
        }}
        body {{
            font-family: 'Noto Sans Devanagari', sans-serif;
            background-color: var(--bg-main);
            color: var(--text-dark);
            margin: 0;
            padding: 25px;
            line-height: 1.8;
            font-size: 18px;
        }}
        .container {{
            max-width: 1200px;
            margin: 0 auto;
        }}
        .header {{
            background: linear-gradient(135deg, #09090b 0%, #1e1b4b 50%, #0f766e 100%);
            color: white;
            padding: 35px 40px;
            border-radius: 20px;
            box-shadow: 0 15px 30px -5px rgba(0,0,0,0.25);
            margin-bottom: 35px;
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 20px;
        }}
        .header h1 {{ margin: 0 0 10px 0; font-size: 32px; font-weight: 900; }}
        .header p {{ margin: 0; opacity: 0.95; font-size: 19px; }}
        .print-btn {{
            background: #ffffff;
            color: #1e1b4b;
            border: none;
            padding: 14px 24px;
            border-radius: 12px;
            font-weight: 800;
            font-size: 17px;
            cursor: pointer;
            font-family: inherit;
            box-shadow: 0 6px 12px rgba(0,0,0,0.15);
        }}
        .search-bar {{
            width: 100%;
            padding: 15px 22px;
            font-size: 19px;
            border: 2px solid #cbd5e1;
            border-radius: 14px;
            margin-bottom: 30px;
            font-family: inherit;
            box-sizing: border-box;
            background: white;
        }}
        .section-title {{
            font-size: 26px;
            font-weight: 800;
            margin: 40px 0 20px 0;
            padding: 12px 20px;
            background: #ffffff;
            border-radius: 12px;
            border-left: 8px solid var(--primary-light);
            box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05);
            color: #0f172a;
        }}
        .badge {{ padding: 6px 14px; border-radius: 8px; font-size: 15px; font-weight: 800; color: white; display: inline-block; }}
        .badge-pre {{ background-color: var(--accent-gold); }}
        .badge-raj {{ background-color: var(--accent-purple); }}
        .badge-5m {{ background-color: #0284c7; }}
        .badge-10m {{ background-color: #4338ca; }}
        
        .fact-card, .mains-card {{
            background: var(--card-bg);
            border-radius: 16px;
            padding: 28px;
            margin-bottom: 25px;
            box-shadow: 0 6px 12px -2px rgba(0,0,0,0.08);
            border: 1px solid var(--border-color);
        }}
        .fact-card {{ border-left: 8px solid var(--accent-gold); }}
        .mains-card {{ border-left: 8px solid #4338ca; }}
        
        .card-header, .mains-header {{ display: flex; align-items: center; gap: 14px; margin-bottom: 14px; }}
        .topic-title {{ font-weight: 800; font-size: 20px; color: #0f172a; }}
        .fact-text {{ margin: 0; color: var(--text-dark); font-size: 19px; line-height: 1.8; }}
        .question-text {{ margin: 12px 0 16px 0; color: #1e1b4b; font-size: 22px; font-weight: 800; }}
        .answer-box {{ background: #f8fafc; padding: 22px; border-radius: 12px; font-size: 18px; border: 1px solid #e2e8f0; }}
        .ans-label {{ font-size: 19px; color: #1e3a8a; display: inline-block; margin-bottom: 6px; }}
        .ans-section {{ margin-bottom: 16px; line-height: 1.8; }}
        .subject-tag {{ color: #475569; font-weight: 700; margin-left: auto; font-size: 16px; }}

        @media print {{
            .print-btn, .search-bar {{ display: none; }}
            body {{ background: white; padding: 0; font-size: 14pt; }}
            .container {{ max-width: 100%; }}
            .fact-card, .mains-card {{ break-inside: avoid; border: 1px solid #999; margin-bottom: 20px; }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <div>
                <h1>📚 RPSC RAS 5-Month Master Revision Notes</h1>
                <p>विगत 5 माह के प्रमुख टेलीग्राम स्टडी चैनल्स से संकलित शुद्ध, प्रामाणिक व परीक्षा-उपयोगी नोट्स (Pre & Mains)</p>
            </div>
            <div>
                <button class="print-btn" onclick="window.print()">🖨️ PDF प्रिंट करें</button>
            </div>
        </div>

        <input type="text" id="searchInput" class="search-bar" placeholder="🔍 किसी भी विषय, हवेली, नदी, चित्रकला, प्रजामण्डल या कीवर्ड से खोजें..." onkeyup="filterNotes()">

        <div class="section-title">🟡 प्रारंभिक परीक्षा तथ्य (RAS & UPSC Prelims Tracker - 5 माह का संपूर्ण संकलन)</div>
        <div id="prelimsContainer">
            {prelims_html}
        </div>

        <div class="section-title">🔵 मुख्य परीक्षा मॉडल उत्तर लेखन सेट (RAS Mains - 5M एवं 10M विशेष प्रश्नोत्तर)</div>
        <div id="mainsContainer">
            {mains_html}
        </div>
    </div>

    <script>
        function filterNotes() {{
            let q = document.getElementById('searchInput').value.toLowerCase();
            let cards = document.querySelectorAll('.fact-card, .mains-card');
            cards.forEach(card => {{
                let text = card.innerText.toLowerCase();
                card.style.display = text.includes(q) ? 'block' : 'none';
            }});
        }}
    </script>
</body>
</html>
"""
    output_dir = os.path.join(os.path.dirname(__file__), "Output_Notes")
    os.makedirs(output_dir, exist_ok=True)
    
    html_path = os.path.join(output_dir, "Rajasthan_5_Months_Telegram_Master_Question_Bank.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    print(f"✅ Clean Notes HTML saved: {html_path}")

    # Build matching clean Word Document
    doc = Document()
    heading = doc.add_heading("RPSC RAS - 5 Month Master Study Notes", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("विगत 5 माह के टेलीग्राम चैनल्स से संकलित शुद्ध, प्रामाणिक मास्टर नोट्स (Pre & Mains Pattern)\n" + "="*60)

    doc.add_heading("1. प्रारंभिक परीक्षा तथ्य (RAS Prelims Tracker)", level=1)
    for it in prelims_facts:
        p = doc.add_paragraph()
        p.add_run(f"• [{it['topic']}] ").bold = True
        p.add_run(it['fact'])

    doc.add_heading("2. मुख्य परीक्षा मॉडल उत्तर सेट (RAS Mains - 5M & 10M)", level=1)
    for q in mains_questions:
        p = doc.add_paragraph()
        p.add_run(f"प्रश्न [{q['marks']} अंक - {q['paper']}]: {q['question']}\n").bold = True
        if q['marks'] == 5:
            p.add_run(f"उत्तर:\n{q['model_answer'].replace('<strong>', '').replace('</strong>', '')}\n")
        else:
            p.add_run(f"भूमिका: {q['intro']}\n")
            p.add_run(f"मुख्य भाग: {q['body'].replace('<strong>', '').replace('</strong>', '').replace('<br>', '\n')}\n")
            p.add_run(f"निष्कर्ष: {q['conclusion']}\n")
        doc.add_paragraph("-" * 40)

    docx_path = os.path.join(output_dir, "Rajasthan_5_Months_Telegram_Master_Question_Bank.docx")
    doc.save(docx_path)
    print(f"✅ Clean Notes DOCX saved: {docx_path}")

    # Sync to Google Drive Desktop Folder
    uploader = DriveSyncUploader()
    uploader.sync_to_drive(html_path)
    uploader.sync_to_drive(docx_path)
    print("✅ Synced clean notes to Google Drive desktop folder!")

if __name__ == '__main__':
    create_master_notes()
